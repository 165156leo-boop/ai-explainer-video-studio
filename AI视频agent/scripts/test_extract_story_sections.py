import tempfile
import unittest
from pathlib import Path

from docx import Document
from extract_story_sections import (
    extract,
    infer_speaker,
    split_dialogue_visual_note,
    split_stage_direction,
)


class ExtractionTest(unittest.TestCase):
    def test_real_docx_preserves_dialogue_and_excludes_visual_note(self):
        with tempfile.TemporaryDirectory(prefix="workbench-extract-test-") as directory:
            source = Path(directory) / "测试.docx"
            document = Document()
            for text in ["5s-25s | 提示一", "工作人员：请核对这张发票。（递出发票）", "审计员：好的。"]:
                document.add_paragraph(text)
            document.save(source)
            result = extract(source)
            dialogues = [item for section in result["sections"] for example in section["examples"]
                         for item in example["items"] if item["type"] == "dialogue"]
            self.assertEqual([item["text"] for item in dialogues], ["请核对这张发票。", "好的。"])
            self.assertEqual(split_dialogue_visual_note("请核对。（递出发票）"), ("请核对。", "递出发票"))

    def test_excludes_postproduction_text_but_keeps_visual_effects(self):
        with tempfile.TemporaryDirectory(prefix="workbench-overlay-test-") as directory:
            source = Path(directory) / "公务接待.docx"
            document = Document()
            for text in [
                "5s-25s | 提示一：接待事项要真实",
                "场景设定：办公室内，审计员与老师核对处理笺。",
                "审计员：请核对处理笺信息。",
                "镜头/特效：",
                "处理笺扫描件浮现，错误信息被红框圈出。",
                "特效：红框轻微闪烁后定格。",
                "弹出绿色提示框：✅ 信息必须真实准确",
                "字幕（屏幕文字）：",
                "❌ 填报信息与实际不符",
                "✅ 正确做法：及时如实填报",
            ]:
                document.add_paragraph(text)
            document.save(source)

            result = extract(source)
            items = result["sections"][0]["examples"][0]["items"]
            item_texts = [item["text"] for item in items]
            excluded_texts = [item["text"] for item in result["excluded_paragraphs"]]

            self.assertIn("处理笺扫描件浮现，错误信息被红框圈出。", item_texts)
            self.assertIn("特效：红框轻微闪烁后定格。", item_texts)
            self.assertNotIn("镜头/特效：", item_texts)
            self.assertIn("镜头/特效：", excluded_texts)
            self.assertIn("弹出绿色提示框：✅ 信息必须真实准确", excluded_texts)
            self.assertIn("字幕（屏幕文字）：", excluded_texts)
            self.assertIn("❌ 填报信息与实际不符", excluded_texts)
            self.assertIn("✅ 正确做法：及时如实填报", excluded_texts)

            # 回归：真实事故里「镜头/特效：」下面的正文被当成现场动作，
            # 一路进了分镜动作和画面词，最后被生图模型画成一张餐桌。
            # 正文必须是 visual_effect，而不是 action_or_narration。
            for text in ("处理笺扫描件浮现，错误信息被红框圈出。", "特效：红框轻微闪烁后定格。"):
                matched = [item for item in items if item["text"] == text]
                self.assertEqual([item["type"] for item in matched], ["visual_effect"])
            self.assertEqual(
                [item for item in items if item["type"] == "action_or_narration"],
                [],
                "特效块正文不能落成现场动作",
            )

    def test_effect_block_closes_at_prompt_box(self):
        """提示框是另一套后期构造，应终止特效块，而不是被一起吞成特效。"""
        with tempfile.TemporaryDirectory(prefix="workbench-effect-scope-") as directory:
            source = Path(directory) / "公务接待.docx"
            document = Document()
            for text in [
                "5s-25s | 提示一：接待事项要真实",
                "场景设定：办公室里，审计小陈站在办公桌前，与王老师沟通。",
                "小陈：王老师，学院在三年前报销了几笔公务接待费用。",
                "镜头/特效：",
                "处理笺扫描件浮现在餐桌上方，来访人员信息被红框圈出。",
                "弹出绿色提示框：✅ 公务接待处理笺信息必须真实准确",
                "小陈低头翻看处理笺，在纸上圈出两处信息。",
            ]:
                document.add_paragraph(text)
            document.save(source)

            result = extract(source)
            items = result["sections"][0]["examples"][0]["items"]
            effects = [item["text"] for item in items if item["type"] == "visual_effect"]
            self.assertEqual(effects, ["处理笺扫描件浮现在餐桌上方，来访人员信息被红框圈出。"])
            self.assertTrue(
                any(
                    item["type"] == "action_or_narration" and "低头翻看处理笺" in item["text"]
                    for item in items
                ),
                "提示框之后的正常画面动作必须恢复为现场动作",
            )

            # 特效提到「餐桌」，本节场景却是「办公室」，必须给出复核旗标。
            flags = result["review_flags"]
            self.assertTrue(
                any("餐桌" in flag and "办公室" in flag for flag in flags),
                f"应提示特效与本节场景不一致，实际旗标：{flags}",
            )

    def test_stage_direction_is_preserved_as_performance_cue(self):
        """回归：真实 Word 里对白标签本身带现场提示，旧版把它们整段剥掉。

        剥掉的后果不止是丢信息：对白镜因此没有任何可执行动作，下游
        `dialogueAction` 只能凭空合成「小陈对王老师说明核查发现」这类口播意向串，
        再被 `continuityVisualActions` 的正则全部删掉 → actions 为空 →
        `normalizeVideoPromptBriefs` 抛 AiInputRejected，画面词与视频词一起卡死。
        """
        with tempfile.TemporaryDirectory(prefix="workbench-cue-") as directory:
            source = Path(directory) / "公务接待.docx"
            document = Document()
            for text in [
                "5s-25s | 提示一：接待事项要真实",
                "王老师一愣：“我们接待时没做好记录，也没及时报销，可能事后把信息填错了。”",
                "小陈温和地提醒：“这里有两处信息需要核对。”",
                "小李走进来，把平板递过去：“这是上个月的接待审批单。”",
                "小赵现场对话学院相关人员：“这几笔费用的用途需要说明。”",
                "刘主任有些为难：“当时确实没按流程走。”",
            ]:
                document.add_paragraph(text)
            document.save(source)

            result = extract(source)
            items = result["sections"][0]["examples"][0]["items"]
            cues = {item["speaker"]: item.get("performance_cue") for item in items}

            # 说话人识别结果不变——提示是被分流，不是被换掉。
            self.assertEqual(
                list(cues),
                ["王老师", "小陈", "小李", "小赵", "刘主任"],
            )
            self.assertEqual(cues["王老师"], "一愣")
            self.assertEqual(cues["小陈"], "温和地提醒")
            self.assertEqual(cues["小李"], "走进来，把平板递过去")
            self.assertEqual(cues["小赵"], "现场对话学院相关人员")
            self.assertEqual(cues["刘主任"], "有些为难")
            # 台词正文不受影响。
            self.assertTrue(
                any(item["text"].startswith("我们接待时没做好记录") for item in items)
            )

    def test_stage_direction_does_not_damage_plain_or_parenthetical_labels(self):
        """提示切分必须对无提示标签与括号型标签保持沉默。"""
        # 无提示：说话人原样保留，不产生空提示。
        self.assertEqual(split_stage_direction("小陈"), ("小陈", None))
        self.assertEqual(split_stage_direction("刘老师"), ("刘老师", None))
        # 括号型：「审计人员A（站在门口）」的括号是场景说明，正文不是提示，
        # 不能把「（」留在说话人里、也不能把「站在门口）」当成提示。
        self.assertEqual(split_stage_direction("审计人员A（站在门口）"), ("审计人员A（站在门口）", None))
        self.assertEqual(
            infer_speaker("审计人员A（站在门口）"),
            ("审计人员A", False, "站在门口", None),
        )
        # 名字里恰好含提示字样的说话人（赵老师）不得被误切。
        self.assertEqual(infer_speaker("赵老师"), ("赵老师", False, None, None))

    def test_estimate_speech_uses_the_workbench_speech_rate(self):
        """语速基线必须与工作台 shot-planning.ts 的 5.5 字/秒 对齐。"""
        from extract_story_sections import estimate_speech

        # 17 个口播单元：6.0 字/秒 会得到 2.8，5.5 字/秒 得到 3.1。
        self.assertEqual(
            estimate_speech("这是一句用于核对语速基线的测试台词")["estimated_seconds_fast"],
            3.1,
        )
        # 下限 2.0 秒，避免极短台词算出 0.4 秒这种不可用值。
        self.assertEqual(estimate_speech("好的")["estimated_seconds_fast"], 2.0)


if __name__ == "__main__":
    unittest.main()
